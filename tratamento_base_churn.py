from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# -----------------------------------------------------------------------------
# 1. Configurações
# -----------------------------------------------------------------------------

ARQUIVO_ENTRADA = Path("base_churn_5000_VF2_aluno.xlsx")
ARQUIVO_SAIDA = Path("base_churn_nova_tratada.xlsx")
DATA_CORTE = pd.Timestamp("2026-09-21")
DIAS_PREVISAO = 60
FIM_JANELA = DATA_CORTE + pd.Timedelta(days=DIAS_PREVISAO)

# -----------------------------------------------------------------------------
# 2. Funções auxiliares
# -----------------------------------------------------------------------------


def converter_data(valor):
    """Converte datas brasileiras e datas do Excel sem interromper o programa."""
    return pd.to_datetime(valor, dayfirst=True, errors="coerce")


def padronizar_contrato(valor):
    """Converte grafias diferentes para o domínio oficial de contratos."""
    if pd.isna(valor):
        return pd.NA

    valor = str(valor).strip().lower()
    mapa = {
        "mensal": "Mensal",
        "anual": "Anual",
        "bienal": "Bienal",
    }
    return mapa.get(valor, pd.NA)


def padronizar_pagamento(valor):
    """Converte grafias diferentes para o domínio oficial de pagamentos."""
    if pd.isna(valor):
        return pd.NA

    valor = str(valor).strip().lower()
    mapa = {
        "boleto": "Boleto",
        "cartão": "Cartão",
        "cartao": "Cartão",
        "pix": "Pix",
    }
    return mapa.get(valor, pd.NA)


def contar_chamados_90_dias(valor):
    """Conta chamados ocorridos nos 90 dias anteriores à data de corte."""
    if pd.isna(valor) or not str(valor).strip():
        return 0

    quantidade = 0
    inicio_janela = DATA_CORTE - pd.Timedelta(days=90)

    for texto_data in str(valor).split(","):
        data = converter_data(texto_data.strip())
        if pd.notna(data) and inicio_janela < data <= DATA_CORTE:
            quantidade += 1

    return quantidade


def meses_entre(data_inicio, data_fim):
    """Calcula o tempo completo de relacionamento em meses."""
    if pd.isna(data_inicio):
        return pd.NA

    meses = (data_fim.year - data_inicio.year) * 12
    meses += data_fim.month - data_inicio.month

    if data_fim.day < data_inicio.day:
        meses -= 1

    return max(0, meses)


# -----------------------------------------------------------------------------
# 3. Leitura e preparação da base original
# -----------------------------------------------------------------------------

if not ARQUIVO_ENTRADA.exists():
    raise FileNotFoundError(
        f"Arquivo não encontrado: {ARQUIVO_ENTRADA.resolve()}"
    )

base = pd.read_excel(ARQUIVO_ENTRADA, dtype={"id_cliente": "string"})
base.columns = [col.strip() for col in base.columns]

# Converte as datas para permitir validações temporais.
# A aplicação célula a célula aceita tanto o padrão brasileiro
# DD/MM/AAAA quanto o padrão ISO AAAA-MM-DD usado em parte da planilha.
base["data_aquisicao"] = base["data_aquisicao_cliente"].apply(converter_data)
base["data_churn"] = base["data_ocorrencia_churn"].apply(converter_data)

# Guarda os valores originais antes de qualquer correção.
base_original = base.copy()

# -----------------------------------------------------------------------------
# 4. Padronização dos campos categóricos
# -----------------------------------------------------------------------------

base["tipo_contrato_padronizado"] = base["tipo_contrato"].apply(
    padronizar_contrato
)
base["forma_pagamento_padronizada"] = base["forma_pagamento"].apply(
    padronizar_pagamento
)

# Ausências categóricas são explicitadas na base final.
base["contrato_ausente"] = base["tipo_contrato_padronizado"].isna().astype(int)
base["pagamento_ausente"] = base["forma_pagamento_padronizada"].isna().astype(int)
base["tipo_contrato_padronizado"] = base["tipo_contrato_padronizado"].fillna(
    "Desconhecido"
)
base["forma_pagamento_padronizada"] = base[
    "forma_pagamento_padronizada"
].fillna("Desconhecido")

# -----------------------------------------------------------------------------
# 5. Validação e tratamento das variáveis numéricas
# -----------------------------------------------------------------------------

# Mensalidade válida: de R$ 10,00 até R$ 500,00.
base["mensalidade_valida"] = base["mensalidade_rs"].where(
    base["mensalidade_rs"].between(10, 500)
)
base["mensalidade_ausente"] = base["mensalidade_valida"].isna().astype(int)

# Atrasos válidos: de 0 até 6 parcelas.
base["atrasos_validos"] = base["atrasos_6m"].where(
    base["atrasos_6m"].between(0, 6)
)
base["atrasos_ausente"] = base["atrasos_validos"].isna().astype(int)

# Mediana por tipo de contrato. O fallback global evita novos nulos.
mediana_mensalidade_global = base["mensalidade_valida"].median()
mediana_atrasos_global = base["atrasos_validos"].median()

mediana_mensalidade_contrato = base.groupby(
    "tipo_contrato_padronizado"
)["mensalidade_valida"].transform("median")
mediana_mensalidade_contrato = mediana_mensalidade_contrato.fillna(
    mediana_mensalidade_global
)

mediana_atrasos_contrato = base.groupby("tipo_contrato_padronizado")[
    "atrasos_validos"
].transform("median")
mediana_atrasos_contrato = mediana_atrasos_contrato.fillna(mediana_atrasos_global)

base["mensalidade_tratada"] = base["mensalidade_valida"].fillna(
    mediana_mensalidade_contrato
).round(2)
base["atrasos_tratados"] = base["atrasos_validos"].fillna(
    mediana_atrasos_contrato
).round().astype(int)

# -----------------------------------------------------------------------------
# 6. Remoção de duplicidades
# -----------------------------------------------------------------------------

# Mantém somente a primeira ocorrência de cada cliente.
# A auditoria registra quantas linhas foram removidas.
quantidade_antes = len(base)
base = base.drop_duplicates(subset=["id_cliente"], keep="first").copy()
quantidade_duplicidades = quantidade_antes - len(base)

# -----------------------------------------------------------------------------
# 7. Seleção da fotografia de clientes ativos na data de corte
# -----------------------------------------------------------------------------

clientes_ativos = base[
    (base["data_aquisicao"] <= DATA_CORTE)
    & (base["data_churn"].isna() | (base["data_churn"] > DATA_CORTE))
].copy()

# -----------------------------------------------------------------------------
# 8. Construção da tabela final de modelagem
# -----------------------------------------------------------------------------

clientes = pd.DataFrame()
clientes["id_cliente"] = clientes_ativos["id_cliente"]
clientes["data_corte"] = DATA_CORTE.date()
clientes["data_aquisicao_cliente"] = clientes_ativos["data_aquisicao"].dt.date
clientes["data_churn"] = clientes_ativos["data_churn"].dt.date
clientes["tempo_cliente_meses"] = clientes_ativos["data_aquisicao"].apply(
    lambda data: meses_entre(data, DATA_CORTE)
)
clientes["mensalidade_rs"] = clientes_ativos["mensalidade_tratada"]
clientes["mensalidade_ausente"] = clientes_ativos["mensalidade_ausente"]
clientes["chamados_suporte"] = clientes_ativos["data_chamado"].apply(
    contar_chamados_90_dias
)
clientes["atrasos_6m"] = clientes_ativos["atrasos_tratados"]
clientes["atrasos_ausente"] = clientes_ativos["atrasos_ausente"]
clientes["tipo_contrato"] = clientes_ativos["tipo_contrato_padronizado"]
clientes["forma_pagamento"] = clientes_ativos["forma_pagamento_padronizada"]
clientes["contrato_ausente"] = clientes_ativos["contrato_ausente"]
clientes["pagamento_ausente"] = clientes_ativos["pagamento_ausente"]
clientes["fibra_optica"] = clientes_ativos["fibra_optica"].astype(int)
clientes["debito_automatico"] = clientes_ativos["debito_automatico"].astype(int)

# Churn = 1 se o cancelamento ocorrer nos próximos 60 dias.
clientes["churn"] = (
    clientes["data_churn"].notna()
    & (pd.to_datetime(clientes["data_churn"]) > DATA_CORTE)
    & (pd.to_datetime(clientes["data_churn"]) <= FIM_JANELA)
).astype(int)

clientes = clientes.reset_index(drop=True)

# -----------------------------------------------------------------------------
# 9. Normalização dos chamados de suporte
# -----------------------------------------------------------------------------

chamados_lista = []
for _, linha in base.drop_duplicates("id_cliente").iterrows():
    if pd.isna(linha["data_chamado"]):
        continue

    for texto_data in str(linha["data_chamado"]).split(","):
        data = converter_data(texto_data.strip())
        if pd.notna(data):
            chamados_lista.append(
                {
                    "id_cliente": linha["id_cliente"],
                    "data_chamado": data.date(),
                }
            )

chamados = pd.DataFrame(chamados_lista).drop_duplicates()
if not chamados.empty:
    chamados = chamados.sort_values(["id_cliente", "data_chamado"])

# -----------------------------------------------------------------------------
# 10. Criação da auditoria
# -----------------------------------------------------------------------------

auditoria_linhas = []

for _, linha in base_original.iterrows():
    mensalidade_invalida = pd.isna(linha["mensalidade_rs"]) or not (
        10 <= linha["mensalidade_rs"] <= 500
    )
    if mensalidade_invalida:
        auditoria_linhas.append(
            {
                "id_cliente": linha["id_cliente"],
                "campo": "mensalidade_rs",
                "problema": "Ausente ou fora do intervalo R$ 10–500",
                "tratamento": "Imputação pela mediana do tipo de contrato",
                "severidade": "Alta",
            }
        )

    atrasos_invalidos = pd.isna(linha["atrasos_6m"]) or not (
        0 <= linha["atrasos_6m"] <= 6
    )
    if atrasos_invalidos:
        auditoria_linhas.append(
            {
                "id_cliente": linha["id_cliente"],
                "campo": "atrasos_6m",
                "problema": "Ausente ou fora do intervalo 0–6",
                "tratamento": "Imputação pela mediana do tipo de contrato",
                "severidade": "Alta",
            }
        )

    if pd.isna(linha["tipo_contrato"]) or str(linha["tipo_contrato"]).strip() != padronizar_contrato(linha["tipo_contrato"]):
        auditoria_linhas.append(
            {
                "id_cliente": linha["id_cliente"],
                "campo": "tipo_contrato",
                "problema": "Valor não canônico ou ausente",
                "tratamento": "Padronização para o domínio oficial",
                "severidade": "Média",
            }
        )

    if pd.isna(linha["forma_pagamento"]) or str(linha["forma_pagamento"]).strip() != padronizar_pagamento(linha["forma_pagamento"]):
        auditoria_linhas.append(
            {
                "id_cliente": linha["id_cliente"],
                "campo": "forma_pagamento",
                "problema": "Valor não canônico ou ausente",
                "tratamento": "Padronização para o domínio oficial",
                "severidade": "Média",
            }
        )

for cliente, quantidade in base_original["id_cliente"].value_counts().items():
    if quantidade > 1:
        auditoria_linhas.append(
            {
                "id_cliente": cliente,
                "campo": "id_cliente",
                "problema": f"Cliente repetido {quantidade} vezes",
                "tratamento": "Retenção de uma ocorrência",
                "severidade": "Alta",
            }
        )

auditoria = pd.DataFrame(auditoria_linhas)

# -----------------------------------------------------------------------------
# 11. Exportação do novo arquivo Excel
# -----------------------------------------------------------------------------

dominios = pd.DataFrame(
    [
        ["tipo_contrato", "Mensal"],
        ["tipo_contrato", "Anual"],
        ["tipo_contrato", "Bienal"],
        ["tipo_contrato", "Desconhecido"],
        ["forma_pagamento", "Boleto"],
        ["forma_pagamento", "Cartão"],
        ["forma_pagamento", "Pix"],
        ["forma_pagamento", "Desconhecido"],
    ],
    columns=["campo", "valor_permitido"],
)

resumo = pd.DataFrame(
    [
        ["Arquivo original", ARQUIVO_ENTRADA.name],
        ["Data de corte", DATA_CORTE.date()],
        ["Fim da janela de previsão", FIM_JANELA.date()],
        ["Linhas originais", len(base_original)],
        ["Clientes distintos", base_original["id_cliente"].nunique()],
        ["Duplicidades removidas", quantidade_duplicidades],
        ["Clientes na base final", len(clientes)],
        ["Clientes com churn = 1", int(clientes["churn"].sum())],
        ["Clientes com churn = 0", int((clientes["churn"] == 0).sum())],
    ],
    columns=["indicador", "valor"],
)

with pd.ExcelWriter(
    ARQUIVO_SAIDA,
    engine="openpyxl",
    date_format="DD/MM/YYYY",
) as escritor:
    clientes.to_excel(escritor, index=False, sheet_name="clientes_modelagem")
    chamados.to_excel(escritor, index=False, sheet_name="chamados_suporte")
    auditoria.to_excel(escritor, index=False, sheet_name="auditoria")
    dominios.to_excel(escritor, index=False, sheet_name="dominios")
    resumo.to_excel(escritor, index=False, sheet_name="resumo")

# Formatação básica para facilitar a leitura no Excel.
workbook = load_workbook(ARQUIVO_SAIDA)
for planilha in workbook.worksheets:
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions

    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.alignment = Alignment(horizontal="center")

    for coluna in range(1, planilha.max_column + 1):
        valores = [
            str(planilha.cell(linha, coluna).value or "")
            for linha in range(1, min(planilha.max_row, 200) + 1)
        ]
        largura = max(12, max(map(len, valores), default=12) + 2)
        planilha.column_dimensions[get_column_letter(coluna)].width = min(
            largura, 36
        )

workbook.save(ARQUIVO_SAIDA)

print("Base tratada criada com sucesso!")
print(f"Arquivo: {ARQUIVO_SAIDA.resolve()}")
print(f"Clientes na base final: {len(clientes)}")
print(f"Churn = 1: {int(clientes['churn'].sum())}")
print(f"Churn = 0: {int((clientes['churn'] == 0).sum())}")
